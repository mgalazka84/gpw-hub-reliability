from __future__ import annotations
from copy import deepcopy
from pathlib import Path
import json
import re
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

PROJECT = Path(__file__).resolve().parent.parent
BASE = PROJECT / "submission"
FINAL = PROJECT / "final"
EMP = PROJECT / "empirical"
OUT = FINAL / "deliverables"
OUT.mkdir(parents=True,exist_ok=True)
observed = json.loads((EMP / "results/forecast_results_all.json").read_text())
network_summary = pd.read_csv(EMP / "results/network_summary.csv")
features = pd.read_csv(EMP / "results/features_close.csv")
coverage = pd.read_csv(EMP / "results/origin_coverage.csv")
market_nodes = pd.read_csv(EMP / "results/hub_selections.csv.gz")
audit = json.loads((EMP / "source_checks/observed_data_audit.json").read_text())
verification = json.loads((EMP / "results/observed_verification.json").read_text())
assert verification["status"] == "PASS"
primary = observed["close_factor_residual"]
intraday = observed["intraday_factor_residual"]
raw_nodes = market_nodes.loc[(market_nodes.convention=="close") & (market_nodes.network=="raw")]
sd_ratios = raw_nodes.groupby("origin_date").log_return_sd.agg(lambda x:x.max()/x.min())
sd_ratios.to_csv(FINAL / "observed_sd_ratios.csv",header=["max_to_min_log_return_sd"])
stats = pd.read_csv(BASE / "results/simulation_summary.csv").set_index(["case","transform"])
variance = pd.read_csv(BASE / "results/variance_sensitivity_summary.csv")
data_summary = json.loads((BASE / "data/constituents_summary.json").read_text())
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin,section.bottom_margin = Inches(.78),Inches(.72)
section.left_margin,section.right_margin = Inches(.82),Inches(.82)
normal=doc.styles["Normal"]
normal.font.name="Times New Roman"; normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1.08
normal.paragraph_format.space_after=Pt(7)
for name,size in [("Title",22),("Subtitle",11),("Heading 1",14),("Heading 2",12),("Caption",10)]:
    style=doc.styles[name];style.font.name="Times New Roman";style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(7)
    if name.startswith("Heading"):
        style.font.bold=True;style.paragraph_format.space_before=Pt(12)
        style.paragraph_format.keep_with_next=True
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))): border.getparent().remove(border)
doc.core_properties.title="Hub reliability and diversification risk in Polish stock correlation networks"
doc.core_properties.author="Marek Gałązka; Hanna Wdowicka"
doc.core_properties.subject="Analytical diagnostics, Monte Carlo evidence and observed GPW results from 2007 to 2025"
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)

def p(text,bold=False,style=None):
    para=doc.add_paragraph(style=style);r=para.add_run(text);r.bold=bold
    if style == "Caption" and text.startswith("Table"):
        para.paragraph_format.keep_with_next=True
    return para
def h(text):doc.add_heading(text,1)
def h2(text):doc.add_heading(text,2)
def link_paragraph(label,url):
    para=doc.add_paragraph()
    para.add_run(label+' ')
    link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'174A7E');properties.append(color)
    run.append(properties);text=OxmlElement('w:t');text.text=url;run.append(text)
    link.append(run);para._p.append(link)
    return para
def mr(text):
    r=OxmlElement("m:r");t=OxmlElement("m:t");t.text=text;t.set(qn("xml:space"),"preserve");r.append(t);return r
def nodes(v):return [mr(v)] if isinstance(v,str) else (v if isinstance(v,list) else [v])
def slot(tag,v):
    element=OxmlElement("m:"+tag)
    for item in nodes(v):element.append(deepcopy(item))
    return element
def sub(base,index):
    e=OxmlElement("m:sSub");e.append(slot("e",base));e.append(slot("sub",index));return e
def sup(base,index):
    e=OxmlElement("m:sSup");e.append(slot("e",base));e.append(slot("sup",index));return e
def both(base,lower,upper):
    e=OxmlElement("m:sSubSup");e.append(slot("e",base));e.append(slot("sub",lower));e.append(slot("sup",upper));return e
def frac(n,d):
    e=OxmlElement("m:f");e.append(slot("num",n));e.append(slot("den",d));return e
def root(v):
    e=OxmlElement("m:rad");pr=OxmlElement("m:radPr");hidden=OxmlElement("m:degHide");hidden.set(qn("m:val"),"1");pr.append(hidden);e.append(pr);e.append(slot("deg",""));e.append(slot("e",v));return e
def bracket(v):
    e=OxmlElement("m:d");pr=OxmlElement("m:dPr")
    for tag,value in [("begChr","("),("endChr",")")]:
        child=OxmlElement("m:"+tag);child.set(qn("m:val"),value);pr.append(child)
    e.append(pr);e.append(slot("e",v));return e
def summation(index,body,upper=None):
    e=OxmlElement("m:nary");pr=OxmlElement("m:naryPr")
    for tag,val in [("chr","∑"),("limLoc","undOvr"),("supHide","1" if upper is None else "0")]:
        child=OxmlElement("m:"+tag);child.set(qn("m:val"),val);pr.append(child)
    e.append(pr);e.append(slot("sub",index));e.append(slot("sup",upper or ""));e.append(slot("e",body));return e
def hat(v):
    e=OxmlElement("m:acc");pr=OxmlElement("m:accPr");child=OxmlElement("m:chr");child.set(qn("m:val"),"̂");pr.append(child);e.append(pr);e.append(slot("e",v));return e

equation_count=0
def eq(expr):
    global equation_count
    equation_count+=1
    if doc.paragraphs and not doc.paragraphs[-1]._p.findall(qn("m:oMath")):
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(9)
    para.paragraph_format.keep_together=True
    math=OxmlElement("m:oMath")
    for part in expr:math.append(deepcopy(part))
    para._p.append(math);para.add_run(f"    ({equation_count})").font.size=Pt(10)

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    for i,label in enumerate(headers):t.rows[0].cells[i].text=label
    for row in rows:
        cells=t.add_row().cells
        for i,label in enumerate(row):cells[i].text=str(label)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement("w:cantSplit"))
        if ri==0:trpr.append(OxmlElement("w:tblHeader"))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement("w:tcBorders")
            for edge in ["top","bottom","left","right"]:
                b=OxmlElement("w:"+edge);b.set(qn("w:val"),"single");b.set(qn("w:sz"),"4");b.set(qn("w:color"),"AAAAAA");borders.append(b)
            pr.append(borders)
            margins=OxmlElement("w:tcMar")
            for edge,value in [("top","70"),("bottom","70"),("left","90"),("right","90")]:
                m=OxmlElement("w:"+edge);m.set(qn("w:w"),value);m.set(qn("w:type"),"dxa");margins.append(m)
            pr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2);paragraph.paragraph_format.line_spacing=1.0
                paragraph.paragraph_format.keep_with_next=ri<len(t.rows)-1
                for run in paragraph.runs:run.font.size=Pt(10);run.bold=ri==0
    p("").paragraph_format.space_after=Pt(0)
    return t

TITLE="Hub reliability and diversification risk in Polish stock correlation networks"
p(TITLE,style="Title")
p("Marek Gałązka and Hanna Wdowicka",style="Subtitle")
h("Abstract")
abstract=("Can a stock that repeatedly occupies a central network position help predict when diversification will weaken? We examine this question using analytical diagnostics, simulations and observed Polish stock prices. Cross-sectional subtraction imposes exact covariance restrictions, while factor exposure can generate persistent hubs without residual dependence. Joint block resampling and fractional treatment of degree ties define a reliability adjustment to observed hub concentration. Simulations separate common exposure, residual dependence and structure induced by unequal variances. The empirical application covers 155 historical WIG20 and mWIG40 stocks across 231 estimation origins during 2007-2025, with 252-session windows and 500 bootstrap replications. At least one hub reaches a selection frequency of 0.80 in 129 raw-network origins, five demeaned origins and eight WIG-residual origins. Chronological forecasts use 73 complete test windows during 2019-2025. Adding residual-network reliability to financial predictors and observed concentration increases mean squared error by 2.20%; the paired block-bootstrap interval includes zero. An open-to-close sensitivity analysis yields a 2.69% reduction, also inconclusive. Available prices are not certified total returns, membership follows quarterly snapshots, and incomplete future outcomes are excluded. Within this scope, persistent network positions depend strongly on the return representation and do not establish incremental forecasting value.")
p(abstract)
p("Keywords: stock correlation networks; minimum spanning tree; bootstrap stability; factor models; diversification; Polish stock market.")
p("JEL classification: C38, C53, G11, G15.")
h("1. Introduction")
p("A diversified stock portfolio can become much less protective when its constituents move together. Correlation networks make these shared movements visible: highly connected stocks occupy central positions and may appear to summarize market risk. The financial question is whether a position that persists under resampling adds useful information about future diversification, beyond what is already visible in correlations, volatility and network concentration. The minimum spanning tree (MST) introduced into financial analysis by Mantegna (1999) retains N - 1 connections and offers a compact setting for this test.")
p("Gałązka (2011) studied 252 stocks traded in Poland in 2007, comparing an MST with a fully weighted correlation network. The return series were first reduced by their contemporaneous cross-sectional average. This earlier construction motivates a re-examination of how benchmark subtraction affects both signed node strength and hub identity. The question concerns the interpretation of the transformed network; algebraic restrictions do not by themselves overturn an empirical finding about a particular sample.")
p("The analysis builds on established work on financial-network interpretation and estimation. Bonanno et al. (2003) compared real-market trees with random and one-factor models, and Miccichè et al. (2003) investigated degree stability. Eom and Park (2017) connected common-factor filtering to diversification and portfolio performance. Bootstrap validation of MST edges was developed by Tumminello et al. (2007) and Musciotto et al. (2018). Kalyagin et al. (2022) studied uncertainty in identifying a population tree. These studies provide the methodological basis for assessing the conditional meaning of a stable hub.")
p("The financial literature also connects network structure to prediction and portfolio choice. Millington and Niranjan (2020) examined partial-correlation networks and portfolio implications. Wang et al. (2024) evaluated the predictive contribution of connectedness measures in IRFA. Härdle et al. (2026) developed a portfolio-risk decomposition and centrality-constrained allocation using regularized covariance and centrality estimation. Within Poland, Tomeczek (2022) extended Gałązka (2011) with additional centralities and pandemic-period networks. Our focus is the incremental information in resampling reliability after conditioning on the concentration already visible in the estimated tree.")
p("We ask whether conditional hub reliability adds information about subsequent diversification loss after controlling for observed concentration and standard financial predictors. The contribution combines an interpretation check with a nested financial test. Exact covariance restrictions and controlled simulations show why stable hubs can arise under different mechanisms. The observed application then uses the available histories of 155 former or current WIG20/mWIG40 constituents over 2007-2025, while reporting the coverage and limitations of the price and membership data.")
p("The observed result qualifies the apparent prominence of stable hubs. A selection frequency of at least 0.80 occurs in 129 of 231 raw-network origins, compared with eight after WIG-factor residualization. In 73 complete retrospective test windows, adding residual-network reliability raises mean squared forecast error by 2.20%, with uncertainty spanning zero. The intraday estimate has the opposite sign but is also inconclusive. These findings motivate a distinction between a reproducible network position and demonstrated incremental forecasting information; neither the simulations nor the market results identify causal shock transmission.")

h("2. Analytical diagnostics for the return representation")
h2("2.1. Cross-sectional subtraction imposes correlation restrictions")
p("Let r(u) be the N-dimensional vector of returns on observation u, measured over a fixed, complete window. Set M = I - 11′/N and c(u) = Mr(u). The transformation used in the earlier Polish-market study is")
eq([sub("c","i"),mr("(u) = "),sub("r","i"),mr("(u) - "),frac("1","N"),summation("j=1",[sub("r","j"),mr("(u)")],"N")])
p("Proposition 1 (zero-sum restriction). For either the population covariance or the usual sample covariance calculated on this complete panel, Γ = MΣM satisfies Γ1 = 0. If all transformed standard deviations are positive, the correlation matrix R and vector σ of transformed standard deviations satisfy Rσ = 0. Therefore")
eq([mr("Γ1 = 0,     Rσ = 0,     "),summation("j ≠ i",[sub("ρ","ij"),sub("σ","j")]),mr(" = -"),sub("σ","i")])
p("Proof. M1 = 0 and M is symmetric. With D = diag(σ), R = D⁻¹ΓD⁻¹ and D⁻¹σ = 1, so Rσ = D⁻¹Γ1 = 0. Isolating the diagonal term gives the final equality. The argument requires the same observations and the same asset set for every covariance entry. Pairwise deletion or a changing cross-sectional universe can break this exact sample identity.")
p("If the transformed volatilities are equal, every signed off-diagonal row sum is -1. With unequal volatilities, it is the volatility-weighted row sum that is fixed. Thus absolute signed strength, positive strength and degree measure different features. Taking the absolute value of a signed row sum does not recover the amount of bilateral positive dependence. These are standard linear-algebra consequences of the transformation, used here as diagnostics rather than as a claimed new theorem.")
h2("2.2. Independent innovations can produce a structured demeaned network")
p("Suppose the original innovations are independent with covariance diag(v₁, …, vN), and let S = Σᵢvᵢ. Direct multiplication by M gives")
eq([sub("Γ","ij"),mr(" = -"),frac([sub("v","i"),mr(" + "),sub("v","j")],"N"),mr(" + "),frac("S",sup("N","2")),mr("  (i ≠ j)")])
eq([sub("Γ","ii"),mr(" = "),sub("v","i"),bracket([mr("1 - "),frac("2","N")]),mr(" + "),frac("S",sup("N","2"))])
p("Equal variances yield the same off-diagonal correlation, -1/(N - 1), and a non-unique population MST. Unequal variances change both the numerator and the normalizing standard deviations. Low-variance assets can share positive correlations because their benchmark-relative returns contain a substantial common subtraction component. Such a hub describes the dependence of benchmark-relative positions. Its existence does not require bilateral dependence among the original innovations.")
h2("2.3. A common factor can create a population star")
p("Consider rᵢ = βᵢf + εᵢ, with mutually uncorrelated idiosyncratic components uncorrelated with f, positive βᵢ and positive variances. Define the standardized factor exposure aᵢ as")
eq([sub("ρ","ij"),mr(" = "),sub("a","i"),sub("a","j"),mr(",    "),sub("a","i"),mr(" = "),frac([sub("β","i"),sub("σ","f")],root([sup(sub("β","i"),"2"),sup(sub("σ","f"),"2"),mr(" + "),sup(sub("σ","ε,i"),"2")]))])
p("Proposition 2 (factor-only star). If one asset m has a strictly greatest aᵢ, the unique correlation MST is the star centered on m. For each other asset i, the edge (i,m) is the unique highest-correlation edge crossing the singleton cut {i}; the cut property places every such edge in the maximum-correlation spanning tree. Because correlation distance is strictly decreasing in correlation, the same edges form the conventional distance MST. The proof illustrates the established effect of factors on financial-network topology. The hub represents common exposure even when idiosyncratic covariances vanish.")

h("3. Conditional hub reliability and concentration")
h2("3.1. Joint resampling and treatment of ties")
p("For each estimation window we compute three representations: raw log returns, cross-sectionally demeaned log returns and factor-model residuals. In the synthetic experiment the factor is observed directly. In the observed-market application the residual model includes an intercept and a single contemporaneous WIG log return of the same return convention. Historical sector factors were not available in the assembled panel. Coefficients are estimated within each window and refitted in every bootstrap replication. Filtering once and resampling the residuals would omit uncertainty from factor estimation.")
p("The bootstrap draws contiguous blocks of complete daily vectors, including all assets and factor observations. We use circular blocks of length 10, concatenate randomly selected starts and trim the resulting sequence to T observations. The same row indices apply to all columns. This preserves contemporaneous sample dependence and some local temporal dependence, following the rationale of block resampling for dependent observations (Künsch, 1989). Circular wrapping creates artificial block boundaries; sensitivity to block length is required for market data.")
p("In each replicated MST, the hub group contains k = ceiling(0.10N) units of membership. Assets whose degree exceeds the kth-largest degree receive membership one. The remaining membership is shared equally across assets tied at the boundary. This avoids a deterministic ticker-order preference and makes the membership sum exactly k. Exact correlation ties are handled by random relabeling before tree extraction; this convention is not uniform sampling over all admissible tied trees.")
eq([sub("p","i,t"),mr(" = "),frac("1","B"),summation("b=1",both("q","i,t","(b)"),"B"),mr(",     "),summation("i=1",sub("p","i,t"),"N"),mr(" = k")])
p("The pᵢ,t values are conditional selection frequencies, including fractional membership at ties. They measure stability under the specified resampling and return representation. They are neither posterior probabilities of economically important firms nor calibrated confidence levels for causal transmission. A threshold such as 0.80 defines a diagnostic event, not a hypothesis test with a controlled error rate. Near population ties, discontinuous tree selection is particularly difficult to calibrate.")
h2("3.2. Reliability acts as a discount to observed degree concentration")
p("Let dᵢ,t be the degree in the original-window MST, and let qᵢ,t be its fractional top-k membership. Compare observed hub concentration Cₜ with reliability-weighted degree share Hₜ:")
eq([sub("C","t"),mr(" = "),frac(summation("i=1",[sub("q","i,t"),sub("d","i,t")],"N"),[mr("2"),bracket("N - 1")]),mr(",    "),sub("H","t"),mr(" = "),frac(summation("i=1",[sub("p","i,t"),sub("d","i,t")],"N"),[mr("2"),bracket("N - 1")])])
p("Proposition 3 (concentration bound). Since every tree degree is at least one, Σᵢdᵢ = 2(N - 1), and 0 ≤ pᵢ ≤ 1 with Σᵢpᵢ = k, we have")
eq([frac("k",[mr("2"),bracket("N - 1")]),mr(" ≤ "),sub("H","t"),mr(" ≤ "),sub("C","t"),mr(" ≤ "),frac("N + k - 2",[mr("2"),bracket("N - 1")])])
p("The middle inequality follows because assigning k units of weight to the largest observed degrees maximizes the weighted sum. The final bound leaves at least N - k units of degree outside the selected group. The difference Aₜ = Cₜ - Hₜ is therefore a nonnegative reliability discount. Adding Hₜ to a model already containing Cₜ is equivalent to adding this discount with the opposite coefficient. The average selection frequency, k/N, carries almost no changing information and must not be used as a market-state feature.")

h("4. Synthetic experimental design")
p("We use N = 60, T = 252, k = 6, 100 independent Monte Carlo samples per data-generating process (DGP), B = 500 bootstrap replications and a circular block length of 10. Five DGPs and three representations produce 1500 network estimates and 750,000 bootstrap trees. The master seed is 20261003; independent sample streams and subsequent bootstrap seeds are saved. Observations are generated in normalized return units and are not calibrated to observed Polish returns.")
p("The Gaussian DGP is xᵢ = βᵢf + ℓᵢg + σᵢeᵢ, with independent unit-variance Gaussian innovations. Its covariance is ββ′ + ℓℓ′ + diag(σ²). The factor-residual representation conditions on f only. For market cases, β is equally spaced from 0.25 to 1.10 and σ from 1.40 to 0.80, except that asset 0 has β₀ = 2.20 and σ₀ = 0.35. These deliberately separated exposures create a clear population hub, making the interpretation problem visible.")
p("Table 0. Synthetic data-generating processes.",style="Caption")
table(["DGP","Parameters and interpretation"],[
    ("Independent, equal variances","β = ℓ = 0; σᵢ = 1. Both raw and correctly filtered population correlations are tied."),
    ("Independent, unequal variances","β = ℓ = 0; σᵢ spans 0.12 to 2.50 geometrically. Original innovations remain independent."),
    ("Market factor only","Market parameters above; ℓ = 0. Idiosyncratic covariance is diagonal."),
    ("Market + residual dependence","Same market; ℓ₁ = 1.80, ℓ₂ to ℓ₂₁ decline from 0.80 to 0.30, others zero. Dependence remains after conditioning on f."),
    ("Market only, SV + t(5)","Unit-variance t(5) innovations, scaled by independent common stochastic volatility; no residual linear covariance.")
],[2.0,4.8])
p("For the last DGP, log variance follows zᵤ = 0.95zᵤ₋₁ + ηᵤ, with ηᵤ Gaussian of standard deviation 0.20 and a stationary initial draw. The scale is exp{(zᵤ - Var(z)/2)/2}, so its squared expectation is one. We discard 500 burn-in observations. Shared stochastic volatility creates nonlinear dependence despite zero residual linear covariances. Factor filtering therefore removes the linear common exposure without eliminating every form of joint risk.")
p("The reported statistics are maximum conditional hub selection frequency, the fraction of samples with a maximum of at least 0.80, observed maximum degree and population-edge recovery. Recovery is reported only when all population pair correlations are distinct, a conservative sufficient condition for uniqueness. Where population ties remain, no arbitrarily selected tree is treated as ground truth. Standard errors summarize variation across independent Monte Carlo samples. Wilson intervals describe the frequency of the diagnostic event, separately from bootstrap selection frequencies.")

h("5. Simulation results")
case_labels={"independent_equal":"Independent, equal variance","independent_unequal":"Independent, unequal variance","market_only":"Market factor only","market_residual":"Market + residual dependence","market_only_sv_t5":"Market only, SV + t(5)"}
cases=list(case_labels)
kinds=["raw","demeaned","factor_residual"]
p("Table 1. Mean maximum conditional hub selection frequency (Monte Carlo standard error). All observations in this section are synthetic.",style="Caption")
rows=[]
for case in cases:
    row=[case_labels[case]]
    for kind in kinds:
        r=stats.loc[(case,kind)];row.append(f"{r.max_selection_mean:.3f} ({r.max_selection_mc_se:.3f})")
    rows.append(row)
table(["DGP","Raw","Demeaned","Factor residual"],rows,[2.6,1.4,1.4,1.4])
p("Under independent equal-variance innovations, maximum selection frequencies average 0.377 for raw returns, 0.324 after demeaning and 0.374 after factor filtering. No sample in these three representations exceeds the 0.80 diagnostic threshold. This does not imply a zero population frequency: for zero events in 100 samples, the upper 95% Wilson bound is approximately 3.7%.")
p("The unequal-variance independent DGP changes the demeaned result sharply. Its mean maximum selection frequency is 0.944 (Monte Carlo standard error 0.0046), and the threshold is reached in 99 of 100 samples. The 95% Wilson interval is 94.6% to 99.8%. The mean maximum observed degree is 9.95, while mean recovery of the unique population edges is only 15.1%. Hub stability and recovery of the entire edge set can therefore differ substantially.")
p("Market-only raw returns have a maximum selection frequency of 1.000 in all 100 samples. Mean maximum degree is 47.05 and population-edge recovery is 79.7%. Demeaning retains a highly stable hub in this DGP, while correctly specified factor residuals reduce mean maximum frequency to 0.380 and yield zero threshold events. When a second dependence factor remains after market filtering, the residual-network maximum averages 0.999 and exceeds the threshold in every sample. The comparison demonstrates conditional sensitivity to residual linear dependence in these models; it does not provide causal identification.")
p("The stochastic-volatility t(5) market-only DGP also has stable raw and demeaned hubs. After filtering the observed market factor, the mean maximum frequency is 0.408 and no sample reaches 0.80. That result concerns the linear correlation network. Common stochastic volatility still connects the residual distributions and is relevant to tail-risk questions outside the present variance-based target.")
para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
para.paragraph_format.keep_with_next=True
para.add_run().add_picture(str(BASE/"results/manuscript_stability.png"),width=Inches(6.65))
p("Figure 1. Mean maximum selection frequencies across 100 samples for each synthetic DGP. Horizontal intervals are ±1.96 Monte Carlo standard errors. They describe Monte Carlo precision and are not confidence intervals for causal hub status. Plotted values are generated directly from the saved replication table by the accompanying Python code.",style="Caption")
p("Table 2. Percentage of synthetic samples with a maximum selection frequency of at least 0.80 (95% Wilson interval, percentage points).",style="Caption")
rows=[]
for case in cases:
    row=[case_labels[case]]
    for kind in kinds:
        r=stats.loc[(case,kind)]
        row.append(f"{100*r.frequency_max_selection_ge_080:.0f}\n[{100*r.frequency_wilson_low:.1f}, {100*r.frequency_wilson_high:.1f}]")
    rows.append(row)
table(["DGP","Raw","Demeaned","Factor residual"],rows,[2.6,1.4,1.4,1.4])
h2("5.1. Exploratory sensitivity to variance contrast")
p("The independent unequal-variance benchmark spans a large ratio of approximately 20.8 between highest and lowest standard deviations. After the initial diagnostic pilot, we added an exploratory comparison at ratios 2, 4, 8 and 16. This addition is explicitly exploratory. Each ratio uses 100 new Gaussian samples and 500 bootstrap replications with the same N, T and block length; common innovations and bootstrap seeds pair the contrasts within each sample.")
rows=[]
for _,r in variance.iterrows():
    rows.append([f"{r.standard_deviation_ratio:.0f}",f"{r.mean_max_selection:.3f} ({r.mc_se_max_selection:.3f})",f"{100*r.frequency_ge_080:.0f}%",f"[{100*r.wilson_low:.1f}%, {100*r.wilson_high:.1f}%]"])
p("Table 3. Variance-contrast sensitivity for demeaned independent innovations.",style="Caption")
table(["SD ratio","Mean maximum (MC SE)","Frequency ≥0.80","95% Wilson interval"],rows,[.8,2.1,1.6,2.3])
p("The threshold event is absent at standard-deviation ratios 2 and 4, occurs in 15% of samples at ratio 8 and in 94% at ratio 16. The result qualifies the strong benchmark finding: large heterogeneity, rather than demeaning alone, drives reliable hubs in this configuration. These contrast levels are diagnostic choices and do not estimate the variance dispersion or event frequencies of the Polish market.")

h("6. Observed data and empirical design")
h2("6.1. Available historical prices and constituent coverage")
p("The empirical universe is reconstructed from the public WIG20 and mWIG40 portfolio archive (GPW Benchmark, 2026). The extraction contains 150 complete quarterly portfolios dated from March 2007 through September 2025, comprising 4500 constituent records and 171 distinct ISINs. Portfolio counts, rank sequences, ISIN check digits and available weight sums have been checked. At each origin the candidate set is the union of the most recent two index portfolios dated strictly before that session. This applies each revision label only from a later session. Extraordinary replacements and exact announcement histories are not reconstructed; membership is therefore a quarterly-snapshot approximation rather than a certified daily point-in-time history.")
p("Observed daily open, high, low, close and volume histories were downloaded from Bankier's public chart service in October 2026 and matched by ISIN (Bankier.pl, 2026). Histories were available for 164 of the 171 identities, including 49 series ending before the study cutoff. The seven unavailable identities remain recorded in the acquisition manifest. The download contains 751,158 daily security records across each series' full available history; 590 records fail positivity, OHLC-range or volume checks and are treated as missing. No price or volume value is imputed or generated. WIG observations provide the market factor and common session calendar, with data capped at 19 September 2025.")
p("An independent check compares 4241 downloaded closing prices with dated quotations in GPW portfolio documents. There are 3950 strict matches and 33 additional matches at the documents' whole-zloty precision. Another 253 ratios are consistent with historical denomination changes, while five discrepancies remain unresolved. This cross-check supports source identification and price-scale consistency but does not certify every corporate action. The delivered price history exhibits denomination adjustments; its cash-dividend adjustment policy is not certified. We consequently report close-to-close price returns as supplied and a separate open-to-close sensitivity analysis, without treating either as a validated dividend-inclusive return panel.")
p("Eligibility is determined solely from candidate membership and the preceding 252 sessions. The same stocks enter both return conventions: all past returns and volumes must be complete; volume must be positive on at least 90% of sessions; zero close-to-close returns may occur on at most 15%; and log-return variances must be positive. At least 20 assets are required. No forward filling or winsorization is used. Large price movements are flagged and retained when the underlying bars satisfy the stated checks. These conditions restrict the application to the available, sufficiently traded large- and mid-cap histories.")
p("The first network origin is 19 March 2007 and subsequent origins are spaced by 20 WIG sessions. The final origin is 14 August 2025, with its outcome ending on 12 September 2025. All 231 origins satisfy the minimum universe size; 155 distinct stocks enter at least one network, with 41-60 stocks per origin and a mean of 51.9. The stock set selected at an origin is fixed for its future outcome window. If any of its future returns is missing, the entire target is missing in both conventions; no stock is retrospectively removed to make that target available. There are 42 such omissions overall and 10 among the 83 test-period origins. The score comparison therefore uses 73 common complete test origins.")
p("Table 4. Observed data coverage and evaluation sample.",style="Caption")
table(["Quantity","Count or period"],[
    ["Historical identities in quarterly portfolios","171"],
    ["Available source histories / unavailable histories","164 / 7"],
    ["Distinct stocks entering the networks","155"],
    ["Network origins and eligible universe","231; 41-60 stocks; mean 51.9"],
    ["Complete / missing future targets, all origins","189 / 42"],
    ["Validation period / retrospective test period","2015-2018 / 2019-2025"],
    ["Complete / missing future targets, test period","73 / 10"],
    ["First / last scored test origin","23 January 2019 / 14 August 2025"],
],[3.75,3.05])
h2("6.2. Financial outcome and forecasting comparisons")
p("The outcome is inverse squared diversification ratio over the next h = 20 sessions, using the covariance of observed simple price returns and origin-date equal weights. Let σ̂p,t,h be the sample volatility of the fixed-weight diagnostic series, and let σ̂i,t,h be each constituent's future-window volatility. Define")
eq([sub("Y","t,h"),mr(" = "),frac(sup(sub(hat("σ"),"p,t,h"),"2"),sup(bracket(summation("i ∈ Vₜ",[sub("w","i,t"),sub(hat("σ"),"i,t,h")])),"2")),mr(",    "),sub("w","i,t"),mr(" = "),frac("1",sub("N","t"))])
p("With complete observations, positive denominator and nonnegative weights, the sample covariance is positive semidefinite and 0 ≤ Yₜ,h ≤ 1. Values closer to one indicate weaker diversification. The series is a covariance-based portfolio diagnostic; transaction costs, turnover and executable portfolio weights have not been studied. A 20-observation window also makes realized covariance noisy, so longer-horizon sensitivity is important. Expected shortfall at a 95% level is not estimated from these short windows.")
p("The financial base model uses the current 20-session and 252-session diversification outcomes, average raw correlation, 20-session and 252-session WIG volatility, the largest raw-correlation eigenvalue share, zero-volume share and asset count. The first extension adds WIG-residual network concentration Cₜ; the second adds Hₜ to that specification. The nested comparison asks whether reliability contributes information beyond visible concentration. Raw and demeaned networks are secondary representations. Two additional benchmarks predict Y using its current 20-session value or a geometrically weighted covariance estimate over the past 252 simple returns, with fixed decay 0.94. The latter uses normalized geometric weights and a correspondingly weighted mean for each asset, then evaluates the same diversification-ratio formula.")
p("Ridge regression uses an expanding training sample and at least 60 fully matured origin labels. Every fit estimates feature means and scales from training observations only. The penalty is selected separately for each specification from {0.01, 0.1, 1, 10, 100} using 2015-2018 validation losses, and is then held fixed for the 2019-2025 test. Model coefficients continue to update as earlier labels mature. At origin t, a past label is available only if its full outcome window has ended. Predictions are clipped to [0,1]. The 20-session spacing prevents overlap of outcome return observations, although estimation windows and model fits remain dependent.")
p("The principal loss comparison is test MSE for the concentration model minus test MSE for concentration plus reliability. Positive differences favor H. All models are scored on the same complete origins. A paired circular-block interval uses six forecast origins per block and 5000 replications, with blocks of three and twelve as sensitivity checks. This interval resamples the realized loss sequence without retraining models and does not incorporate all model-selection or source-data uncertainty. Secondary representation and intraday comparisons are exploratory, without multiplicity adjustment.")
h2("6.3. Retrospective scope and recorded analysis choices")
p("The empirical design and source hashes were recorded before forecast evaluation, after the historical outcomes already existed. This is a retrospective analysis, not prospective preregistration. An earlier broader design contemplated total returns, sector factors and fully reconciled membership. Availability of the public data led to the explicit price-return, single-WIG-factor and quarterly-snapshot implementation reported here. The supplied scripts record these departures. The intraday convention and the 0.94 covariance benchmark were included before inspecting test scores; no specification was selected because it improved the final test.")
p("Each of three representations and two return conventions uses T = 252, B = 500, ten-session bootstrap blocks and a 10% hub group, yielding 693,000 resampled empirical trees. Independent arithmetic checks recompute every current and future Y, every C and H, tree degree sums, total fractional memberships, all test predictions, reported MSE values and source-response hashes. Assertions verify common complete panels, mature training labels, non-overlapping outcomes and the theoretical bounds. Additional summaries of named firms and volatility dispersion are descriptive analyses of the saved estimates, not further predictive specifications.")

h("7. Results on observed GPW prices")
h2("7.1. Hub stability depends on the return representation")
p("Table 5 summarizes the observed rolling networks. For close-to-close returns, mean concentration falls from 0.303 in raw networks to 0.223 after cross-sectional subtraction and 0.231 after WIG residualization. The corresponding reliability shares are 0.243, 0.160 and 0.162. At least one selection frequency reaches 0.80 in 129 of 231 raw-network origins (55.8%), but in only five demeaned origins (2.2%) and eight residual origins (3.5%). The intraday results display the same broad representation dependence. These are descriptive frequencies over overlapping estimation windows, not independent-binomial significance tests.")
p("Table 5. Mean concentration and reliability in observed networks.",style="Caption")
empirical_rows=[]
for convention in ["close","intraday"]:
    for kind,name in [("raw","Raw"),("demeaned","Demeaned"),("factor_residual","WIG residual")]:
        row=network_summary.loc[(network_summary.convention==convention)&(network_summary.network==kind)].iloc[0]
        empirical_rows.append([("Close-to-close" if convention=="close" else "Open-to-close")+" / "+name,
            f"{row.mean_C:.3f}",f"{row.mean_H:.3f}",f"{row.mean_uncertainty_discount:.3f}",f"{int(row.origins_with_selection_at_least_80)}/231"])
table(["Return representation","Mean C","Mean H","Mean C-H","Origins with p ≥ 0.80"],empirical_rows,[2.55,.8,.8,1.0,1.65])
para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
para.add_run().add_picture(str(EMP/"figures/network_concentration.png"),width=Inches(6.45))
p("Figure 2. Observed concentration C and reliability share H in 231 close-to-close estimation windows. Every point uses the past 252 sessions. Changes reflect both dependence dynamics and changes in the eligible historical universe.",style="Caption")
p("The simulation with a standard-deviation ratio of 20.8 is deliberately severe. In the observed close-to-close windows, the ratio of the largest to smallest unfiltered log-return standard deviation has median 3.99, mean 4.18 and range 2.19-10.21. It is therefore inappropriate to extrapolate the simulation's 99% threshold frequency to these market windows. The much lower empirical frequency for demeaned hubs is consistent with the milder observed variance contrasts, although changing factors, non-Gaussian returns and varying universe sizes prevent a direct calibrated attribution.")
p("Table 6 identifies the five strongest raw-network leaders by mean selection frequency among firms observed in at least 60 windows. Each firm's three means use its own identical eligible windows across representations; coverage differs between firms. PKO BP has mean raw selection frequency 0.617 but residual frequency 0.037, while the corresponding values for Pekao are 0.482 and 0.098. These reductions show how a persistent raw position can depend on the market representation. They do not show that the firms cease to be economically important or that a surviving residual hub transmits shocks.")
firm_means=market_nodes.loc[market_nodes.convention=="close"].groupby(["isin","network"]).bootstrap_selection.mean().unstack("network")
firm_counts=raw_nodes.groupby("isin").size()
leaders=firm_means.loc[firm_counts>=60].sort_values("raw",ascending=False).head(5)
names={"PLPKO0000016":"PKO BP","PLPEKAO00016":"Pekao","PLBZ00000044":"BZ WBK / Santander BP","PLBRE0000012":"BRE / mBank","PLKGHM000017":"KGHM"}
leader_rows=[]
for isin,row in leaders.iterrows():
    leader_rows.append([names.get(isin,isin),isin,int(firm_counts.loc[isin]),f"{row['raw']:.3f}",f"{row['demeaned']:.3f}",f"{row['factor_residual']:.3f}"])
p("Table 6. Mean selection frequencies for descriptive raw-network leaders.",style="Caption")
table(["Historical name","ISIN","Windows","Raw","Demeaned","WIG residual"],leader_rows,[1.4,1.45,.75,.8,1.1,1.3])
leaders.join(firm_counts.rename("eligible_windows")).to_csv(FINAL/"descriptive_hub_leaders.csv")
h2("7.2. Incremental forecasting value is not established")
p("Table 7 reports MSE for the 73 common complete test origins from 23 January 2019 to 14 August 2025. The financial base has MSE 0.009995, compared with 0.012657 for persistence and 0.012582 for the fixed-decay covariance benchmark. Adding residual concentration alone gives MSE 0.009886. Adding reliability to that specification gives 0.010103, a 2.20% increase relative to concentration alone. The paired mean loss difference is -0.000217, with a 95% six-origin block interval [-0.000653, 0.000076]. Thus the primary comparison does not establish an incremental predictive benefit for H; the interval also prevents a precise claim of systematic harm.")
p("Table 7. Test MSE multiplied by 1000 for close-to-close price returns.",style="Caption")
forecast_rows=[]
for model,name in [("base","Financial base"),("degree","Base + C"),("reliability","Base + C + H"),("persistence","Current 20-session Y"),("ewma","EWMA covariance, decay 0.94")]:
    forecast_rows.append([name]+[f"{1000*observed['close_'+kind]['mse'][model]:.3f}" for kind in ["raw","demeaned","factor_residual"]])
table(["Forecast model","Raw network","Demeaned network","WIG residual"],forecast_rows,[2.9,1.3,1.3,1.3])
p("All columns use the same test outcomes. The financial-only benchmarks are shared across network representations. For the primary residual comparison, validation selects ridge penalty 10 for the base, concentration and concentration-plus-reliability models. The exact paired intervals using three- and twelve-origin blocks are [-0.000607, 0.000096] and [-0.000639, 0.000075], respectively; both include zero. Raw and demeaned reliability extensions also fail to improve test MSE. Model selection and forecast errors for every origin are retained in the reproduction files.")
h2("7.3. Intraday sensitivity and timing of gains and losses")
p("Open-to-close returns reduce exposure to overnight price adjustments while excluding overnight risk and dividend income. They answer a different financial question, so agreement with the main analysis would not certify total-return robustness. Table 8 reports the same nested comparisons for these returns. The residual concentration model has MSE 0.005824 and the reliability extension 0.005667, a 2.69% reduction. The paired mean difference is 0.000156, with a 95% six-origin interval [-0.000143, 0.000489]. The positive point estimate therefore remains inconclusive. The financial base, persistence and EWMA MSE values are 0.006062, 0.007765 and 0.006958, respectively.")
p("Table 8. Intraday test errors and paired loss intervals. MSE and interval endpoints are multiplied by 1000; positive MSE reductions favor H.",style="Caption")
intra_rows=[]
for kind,name in [("raw","Raw"),("demeaned","Demeaned"),("factor_residual","WIG residual")]:
    row=observed['intraday_'+kind];ci=row['primary_interval']
    intra_rows.append([name,f"{1000*row['mse']['degree']:.3f}",f"{1000*row['mse']['reliability']:.3f}",f"{100*row['relative_mse_reduction_vs_degree']:.2f}%",f"[{1000*ci['ci95_low']:.3f}, {1000*ci['ci95_high']:.3f}]"])
table(["Representation","Base + C","Base + C + H","MSE reduction","95% loss interval"],intra_rows,[1.6,1.1,1.2,1.1,1.8])
para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
para.add_run().add_picture(str(EMP/"figures/cumulative_loss_comparison.png"),width=Inches(6.4))
p("Figure 3. Cumulative paired squared-error difference for the residual networks, concentration minus concentration-plus-reliability. Upward movement favors H. The curves describe the timing of the realized differences and are not separate subperiod tests.",style="Caption")
p("The opposite signs of the two residual point estimates discourage a general claim of financial improvement. The intraday comparison also uses different penalties selected in validation: 10 for concentration and 100 for concentration plus reliability. Its estimate concerns that forecasting procedure rather than an isolated coefficient under a common penalty. The study does not compare all covariance forecasters, alternative horizons or weekly networks; its predictive conclusion is limited to the stated models, window and observed targets.")

h("8. Discussion and conclusion")
p("The analytical and synthetic results distinguish several mechanisms behind persistent hubs. Raw-return centrality can summarize common exposure; cross-sectional subtraction can introduce structured correlations when variances differ; a filtered network can retain linear dependence under a chosen factor model. The observed Polish data add a financial qualification: frequent raw hubs largely disappear after WIG residualization, and conditional reliability does not demonstrate an incremental forecast gain in the main available-data test. Bootstrap stability remains an interpretation and estimation diagnostic whose financial usefulness must be evaluated against explicit benchmarks.")
p("The scope of the observed evidence is deliberately finite. Seven historical identities could not be downloaded, complete-case past filters exclude additional observations, and the quarterly archive does not contain every extraordinary replacement. Although former constituents and ended histories are included, residual availability and survivorship selection remain possible. The downloaded 2026 vintage may contain back-adjustments that were not available in real time. Cash dividends, every corporate action and terminal delisting outcomes are not fully reconciled. The conclusions concern observed price covariance in the assembled historical WIG20/mWIG40 subset, not the whole GPW or a certified investable total-return strategy.")
para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
para.add_run().add_picture(str(EMP/"figures/coverage.png"),width=Inches(6.2))
p("Figure 4. Eligible stocks and completeness of future targets. Forty-two of 231 targets are missing, including ten of 83 test targets. Their omission may be related to suspensions, exits or data quality; the resulting evaluation is conditional on complete outcomes.",style="Caption")
p("Future-target omission can be nonrandom and its effect is unresolved. In particular, the 73-window test may underrepresent difficult trading or exit episodes. The single WIG factor also contains many of the securities under study, creating possible mechanical dependence; it does not isolate every sector or nonlinear common component. Rolling estimation, short realized-covariance windows, finite bootstrap replication and a limited test span further restrict precision. The loss intervals describe the saved test sequence and do not remove these limitations. A broader dividend-inclusive panel, richer factors and independent market replication would test external validity rather than retroactively convert the present conditional estimates into stronger evidence.")
p("Within the available sample, the conclusion is that stable centrality depends on how returns are represented, while its incremental forecasting value remains unestablished. This extends the earlier Polish-market network description by connecting representation, resampling uncertainty and a directly benchmarked diversification outcome. It supports reporting conditional hub reliability alongside observed concentration and data coverage, with financial claims determined by the forecast comparison rather than by the visual prominence of a node.")

h("Data and code availability")
p("The code and derived research outputs are publicly available in the GitHub repository accompanying this article (Gałązka and Wdowicka, 2026), version v1.0.1. The repository contains simulation parameters and saved seeds, replication tables, observed origin-level features, hub and edge estimates, exclusion logs, forecast predictions, selected penalties, source-response hashes and the manuscript-building code. Its documented offline workflow checks the saved derived results and recreates the figures and manuscript; a price-level rerun requires separately obtained source inputs.")
link_paragraph("Repository:","https://github.com/mgalazka84/gpw-hub-reliability")
link_paragraph("Version used for this article:","https://github.com/mgalazka84/gpw-hub-reliability/tree/v1.0.1")
p("Public price histories were accessed through Bankier's chart service; historical index portfolios are supplied by GPW Benchmark. The public repository records source addresses, retrieval dates and hashes, but excludes raw quotation responses, price/return panels and source portfolio documents because public redistribution rights have not been established. These source inputs are retained separately for author verification subject to provider terms. Later downloads can differ from the download vintage identified in the acquisition manifests. The code distinguishes simulated observations from downloaded prices and does not fill missing market values.")
h("Use of AI assistance")
p("ChatGPT (OpenAI) was used to assist with Python coding and data acquisition and extraction.")

h("References")
references=[
"Bankier.pl. (2026, accessed October). Historical company and WIG quotations [Public chart service]. https://www.bankier.pl/inwestowanie/profile/quote.html?symbol=PKOBP",
"Bonanno, G., Caldarelli, G., Lillo, F., & Mantegna, R. N. (2003). Topology of correlation-based minimal spanning trees in real and model markets. Physical Review E, 68, 046130. https://doi.org/10.1103/PhysRevE.68.046130",
"Eom, C., & Park, J. W. (2017). Effects of common factors on stock correlation networks and portfolio diversification. International Review of Financial Analysis, 49, 1-11. https://doi.org/10.1016/j.irfa.2016.11.007",
"Gałązka, M. (2011). Characteristics of the Polish Stock Market correlations. International Review of Financial Analysis, 20(1), 1-5. https://doi.org/10.1016/j.irfa.2010.11.002",
"Gałązka, M., & Wdowicka, H. (2026). GPW hub reliability: Code and derived research outputs (Version v1.0.1) [Computer software]. GitHub. https://github.com/mgalazka84/gpw-hub-reliability/tree/v1.0.1",
"GPW Benchmark. (2026, accessed October). Historyczne portfele indeksów [Historical index portfolios]. https://gpwbenchmark.pl/historyczne-portfele-indeksow",
"Härdle, W. K., Ren, R., Wang, Z., & Wu, W.-B. (2026). A Network View on Portfolio Risk. Journal of Business & Economic Statistics, advance online publication. https://doi.org/10.1080/07350015.2026.2634827",
"Kalyagin, V. A., Koldanov, A. P., & Koldanov, P. A. (2022). Reliability of maximum spanning tree identification in correlation-based market networks. Physica A, 599, 127482. https://doi.org/10.1016/j.physa.2022.127482",
"Künsch, H. R. (1989). The jackknife and the bootstrap for general stationary observations. The Annals of Statistics, 17(3), 1217-1241. https://doi.org/10.1214/aos/1176347265",
"Mantegna, R. N. (1999). Hierarchical structure in financial markets. The European Physical Journal B, 11, 193-197. https://doi.org/10.1007/s100510050929",
"Miccichè, S., Bonanno, G., Lillo, F., & Mantegna, R. N. (2003). Degree stability of a minimum spanning tree of price return and volatility. Physica A, 324, 66-73. https://doi.org/10.1016/S0378-4371(03)00002-5",
"Millington, T., & Niranjan, M. (2020). Partial correlation financial networks. Applied Network Science, 5, 11. https://doi.org/10.1007/s41109-020-0251-z",
"Musciotto, F., Marotta, L., Miccichè, S., & Mantegna, R. N. (2018). Bootstrap validation of links of a minimum spanning tree. Physica A, 512, 1032-1043. https://doi.org/10.1016/j.physa.2018.08.020",
"Tomeczek, A. F. (2022). A minimum spanning tree analysis of the Polish stock market. Journal of Economics and Management, 44. https://doi.org/10.22367/jem.2022.44.17",
"Tumminello, M., Coronnello, C., Lillo, F., Miccichè, S., & Mantegna, R. N. (2007). Spanning trees and bootstrap reliability estimation in correlation-based networks. International Journal of Bifurcation and Chaos, 17, 2319-2329. https://doi.org/10.1142/S0218127407018415",
"Wang, G.-J., Chen, Y., Zhu, Y., & Xie, C. (2024). Systemic risk prediction using machine learning: Does network connectedness help prediction? International Review of Financial Analysis, 93, 103147. https://doi.org/10.1016/j.irfa.2024.103147",
]
for reference in references:
    para=p(reference);para.paragraph_format.left_indent=Inches(.18);para.paragraph_format.first_line_indent=Inches(-.18)
    for r in para.runs:r.font.size=Pt(10.5)

destination=OUT/"IRFA_manuscript_completed.docx"
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if "so its squared expectation is one" in run.text:
            run.text=run.text.replace("so its squared expectation is one","so its expected square is one")
    if re.search(r"\bTable [0-8]\b",paragraph.text):
        for run in paragraph.runs:
            run.text=re.sub(r"\bTable ([0-8])\b",lambda m:"Table "+str(int(m.group(1))+1),run.text)
doc.save(destination)
(FINAL/"manuscript_build.json").write_text(json.dumps(dict(title=TITLE,
    abstract_words=len(abstract.split()),native_equations=equation_count,references=len(references),
    text_words=sum(len(paragraph.text.split()) for paragraph in doc.paragraphs)),indent=2))
print(destination)
print((FINAL/"manuscript_build.json").read_text())
